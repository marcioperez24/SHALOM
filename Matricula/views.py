from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required, permission_required
from django.contrib import messages
from django.db import transaction
from .forms import EstudianteForm, MatriculaForm, InscripcionesFormset
from datetime import date
from django.http import JsonResponse
from configuracion.models import Materias, Nivel, Seccion, TipoMatricula, Modalidades, Turno, Semestres, Cortes
from Matricula.models import Inscripciones,Estudiantes_Nueva, Matricula,Calificaciones
from django.core.paginator import Paginator
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from django.utils import timezone
from django.views.decorators.http import require_http_methods
from collections import OrderedDict
from decimal import Decimal, ROUND_HALF_UP
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
import openpyxl
from django.db.models import Q

@login_required(login_url='/login/')
@permission_required('Matricula.ver_modulo_Registro', raise_exception=True)
def Matricula_home(request):
    return render(request, 'Matricula/Matricula_home.html', {})

@login_required(login_url='/login/')
@permission_required('Matricula.crear_estudiantes_nueva', raise_exception=True)
def crear_estudiante(request):
    if request.method == 'POST':
        estudiante_form = EstudianteForm(request.POST, request.FILES)
        
        if estudiante_form.is_valid():
            try:
                with transaction.atomic():
                    estudiante = estudiante_form.save(commit=False)
                    estudiante.UsuarioID = request.user
                    estudiante.save()
                    
                    messages.success(request, 'Estudiante registrado exitosamente. Ahora puede continuar con la matrícula.')
                    
                    # Redirigir a una página para el siguiente paso, pasando el ID del estudiante
                    return redirect('Matricula:crear_matricula', estudiante_id=estudiante.EstudianteID)

            except Exception as e:
                messages.error(request, f'Ocurrió un error al guardar el estudiante: {e}')
        else:
            messages.error(request, 'Ocurrió un error en la validación del formulario del estudiante.')
    
    else:  # Petición GET
        estudiante_form = EstudianteForm()
        
    context = {
        'estudiante_form': estudiante_form,
    }
    return render(request, 'Matricula/crear_estudiante.html', context)


@login_required(login_url='/login/')
def get_secciones_por_nivel(request):
    nivel_id = request.GET.get('nivel_id')
    if nivel_id:
        try:
            nivel_id_int = int(nivel_id)
            secciones = Seccion.objects.filter(NivelID=nivel_id_int, Activo=True)
            data = [{'id': seccion.SeccionID, 'nombre': seccion.NombreSeccion} for seccion in secciones]
            return JsonResponse(data, safe=False)
        except (ValueError, Nivel.DoesNotExist):
            return JsonResponse([], safe=False)
    return JsonResponse([], safe=False)

@login_required(login_url='/login/')
@permission_required('Matricula.crear_matricula', raise_exception=True)
def crear_matricula(request, estudiante_id):
    # 1. Obtener el estudiante
    estudiante = get_object_or_404(Estudiantes_Nueva, EstudianteID=estudiante_id)
    
    if request.method == 'POST':
        form = MatriculaForm(request.POST)
        if form.is_valid():
            try:
                with transaction.atomic():
                    # --- LOGICA DE DESACTIVACIÓN ---
                    # Desactivamos matrículas anteriores para que solo haya una activa
                    Matricula.objects.filter(EstudianteID=estudiante, activo=True).update(activo=False)

                    # 2. Guardar la nueva Matrícula
                    nueva_matricula = form.save(commit=False)
                    nueva_matricula.EstudianteID = estudiante
                    nueva_matricula.UsuarioID = request.user
                    nueva_matricula.activo = True 
                    nueva_matricula.save()
                    
                    # 3. Buscar materias del nivel seleccionado para inscripción automática
                    materias_del_nivel = Materias.objects.filter(
                        NivelID=nueva_matricula.NivelID, 
                        activo=True
                    )
                    
                    if not materias_del_nivel.exists():
                        messages.warning(request, f"Matrícula creada, pero no se encontraron materias para el nivel seleccionado.")
                    else:
                        # 4. Inscribir automáticamente al estudiante llenando TODOS los campos
                        # para evitar los NULLs vistos en la base de datos
                        inscripciones_creadas = []
                        for materia in materias_del_nivel:
                            inscripciones_creadas.append(
                                Inscripciones(
                                    MatriculaID=nueva_matricula,
                                    EstudianteID=estudiante,
                                    MateriaID=materia,
                                    # --- NUEVOS CAMPOS PARA EVITAR NULLS ---
                                    NivelID=nueva_matricula.NivelID,
                                    SeccionID=nueva_matricula.SeccionID,
                                    ModalidadID=nueva_matricula.ModalidadID,
                                    TurnoID=nueva_matricula.TurnoID,
                                    AñoMatricula=nueva_matricula.AñoMatricula
                                )
                            )
                        
                        # Guardado masivo de inscripciones con todos sus datos
                        Inscripciones.objects.bulk_create(inscripciones_creadas)
                        
                        messages.success(request, f"Matrícula y {len(inscripciones_creadas)} materias registradas correctamente.")
                    
                    return redirect('Matricula:listar_estudiantes')
            
            except Exception as e:
                messages.error(request, f"Error al procesar la matrícula: {str(e)}")
    else:
        form = MatriculaForm()
    
    return render(request, 'Matricula/crear_matricula.html', {
        'matricula_form': form,
        'estudiante': estudiante
    })

@login_required(login_url='/login/')
def crear_inscripciones(request, matricula_id):
    matricula = get_object_or_404(Matricula, pk=matricula_id)
    estudiante = matricula.EstudianteID

    inscripciones_queryset = Inscripciones.objects.filter(MatriculaID=matricula)
    
    if request.method == 'POST':
        formset = InscripcionesFormset(request.POST, prefix='inscripciones', 
                                       queryset=inscripciones_queryset,
                                       form_kwargs={'matricula': matricula})
        if formset.is_valid():
            try:
                with transaction.atomic():
                    # Manejar la lógica de guardado aquí...
                    for form in formset:
                        if form.has_changed() and not form.cleaned_data.get('DELETE', False):
                            inscripcion = form.save(commit=False)
                            inscripcion.MatriculaID = matricula
                            inscripcion.EstudianteID = estudiante
                            inscripcion.NivelID = matricula.NivelID
                            inscripcion.SeccionID = matricula.SeccionID
                            inscripcion.ModalidadID = matricula.ModalidadID
                            inscripcion.TurnoID = matricula.TurnoID
                            inscripcion.AñoMatricula = matricula.AñoMatricula
                            inscripcion.save()
                    
                    # Guardar las inscripciones marcadas para eliminación
                    for form in formset.deleted_forms:
                        form.instance.delete()

                    messages.success(request, 'Inscripción de materias completada exitosamente.')
                    return redirect('Matricula:listar_estudiantes')

            except Exception as e:
                messages.error(request, f'Ocurrió un error al guardar las inscripciones: {e}')
        else:
            messages.error(request, 'Ocurrió un error en la validación de las inscripciones.')
    else:
        # Petición GET: se muestra el formset con las inscripciones existentes
        formset = InscripcionesFormset(prefix='inscripciones', 
                                       queryset=inscripciones_queryset,
                                       form_kwargs={'matricula': matricula})

    materias_del_nivel = Materias.objects.filter(NivelID=matricula.NivelID, activo=True)
    total_materias_nivel = materias_del_nivel.count()
    materias_inscritas_ids = inscripciones_queryset.values_list('MateriaID', flat=True)
    materias_faltantes = materias_del_nivel.exclude(MateriaID__in=materias_inscritas_ids)
    
    materias_inscritas_count = inscripciones_queryset.count()
    todas_inscritas = materias_inscritas_count >= total_materias_nivel and total_materias_nivel > 0

    context = {
        'estudiante': estudiante,
        'matricula': matricula,
        'formset': formset,
        'total_materias_nivel': total_materias_nivel,
        'materias_inscritas_count': materias_inscritas_count,
        'todas_inscritas': todas_inscritas,
        'materias_faltantes': materias_faltantes,
    }
    return render(request, 'Matricula/crear_inscripciones.html', context)


@login_required(login_url='/login/')
def listar_estudiantes(request):
    anio_actual = timezone.now().year
    query = request.GET.get('q', '')
    filtro = request.GET.get('filtro', 'todos')
    nivel_id = request.GET.get('nivel', '')  # Captura el nivel del selector
    ordenamiento = '-FechaRegistro' 
    
    # 1. Obtener lista de niveles para el dropdown del template
    niveles = Nivel.objects.all().order_by('NombreNivel')
    
    # 2. Queryset base: Solo estudiantes activos
    estudiantes_queryset = Estudiantes_Nueva.objects.filter(activo=True)
    
    # 3. Aplicar Filtro de Búsqueda (Nombre o Código)
    if query:
        estudiantes_queryset = estudiantes_queryset.filter(
            Q(NombreCompleto__icontains=query) | Q(Codigo_MINED__icontains=query)
        )

    # 4. Aplicar Filtro por Nivel / Grado
    if nivel_id:
        # Filtramos estudiantes que tengan una matrícula activa en el nivel seleccionado
        estudiantes_queryset = estudiantes_queryset.filter(
            matricula__NivelID_id=nivel_id,
            matricula__activo=True
        ).distinct()

    # 5. Construir lista con datos extendidos (Matrículas)
    lista_estudiantes_completa = []
    total_registrados = estudiantes_queryset.count()
    total_matriculados_actual = 0
    
    # Ordenamos y procesamos cada estudiante
    for est in estudiantes_queryset.order_by(ordenamiento):
        # Buscamos la matrícula activa del estudiante
        matricula_activa = est.matricula_set.filter(activo=True).first()
        # Verificamos si la matrícula pertenece al año escolar actual
        es_periodo_actual = bool(matricula_activa and matricula_activa.AñoMatricula == anio_actual)
        
        if es_periodo_actual:
            total_matriculados_actual += 1

        datos = {
            'estudiante': est,
            'matricula_activa': matricula_activa,
            'es_periodo_actual': es_periodo_actual
        }
        lista_estudiantes_completa.append(datos)

    # 6. Aplicar filtros de estado (Botones/Tarjetas superiores)
    if filtro == 'matriculados':
        lista_final = [e for e in lista_estudiantes_completa if e['es_periodo_actual']]
    elif filtro == 'pendientes':
        lista_final = [e for e in lista_estudiantes_completa if not e['es_periodo_actual']]
    else:
        lista_final = lista_estudiantes_completa

    # 7. Lógica de Paginación (10 por página)
    page_number = request.GET.get("page")
    paginator = Paginator(lista_final, 10)
    
    try:
        page_obj = paginator.page(page_number)
    except (PageNotAnInteger, EmptyPage):
        page_obj = paginator.page(1)
        
    # Rango de páginas para mostrar (ej: [1, 2, 3])
    paginator_range = range(max(1, page_obj.number - 2), min(paginator.num_pages, page_obj.number + 2) + 1)
    
    # Preservar parámetros de búsqueda en los links de paginación
    query_params = request.GET.copy()
    if 'page' in query_params:
        query_params.pop('page')
        
    context = {
        "estudiantes": page_obj, 
        "query": query,
        "filtro_actual": filtro,
        "nivel_seleccionado": nivel_id,
        "niveles": niveles,
        "paginator_range": paginator_range,
        "get_params": query_params.urlencode(),
        "anio_actual": anio_actual,
        "stats": {
            "total": total_registrados,
            "matriculados": total_matriculados_actual,
            "pendientes": total_registrados - total_matriculados_actual
        }
    }
    return render(request, "matricula/listar_estudiantes.html", context)

@login_required(login_url='/login/')
@permission_required('Matricula.desactivar_estudiantes_nueva', raise_exception=True)
def desactivar_estudiante(request, pk):
    # Buscamos al estudiante por su ID
    estudiante = get_object_or_404(Estudiantes_Nueva, EstudianteID=pk)
    
    # Cambiamos el campo 'activo' a False (según tu diagrama de base de datos)
    estudiante.activo = False
    estudiante.save()
    
    messages.success(request, f"El estudiante {estudiante.NombreCompleto} ha sido dado de baja correctamente.")
    return redirect('Matricula:listar_estudiantes')


@login_required(login_url='/login/')
@permission_required('Matricula.ver_inactivos_estudiantes_nueva', raise_exception=True)
def listar_estudiantes_inactivos(request):
    anio_actual = timezone.now().year
    query = request.GET.get('q', '')
    
    # Filtrar solo los inactivos
    estudiantes_queryset = Estudiantes_Nueva.objects.filter(activo=False)
    
    if query:
        estudiantes_queryset = estudiantes_queryset.filter(
            Q(NombreCompleto__icontains=query) | Q(Codigo_MINED__icontains=query)
        )

    lista_inactivos_completa = []
    
    # Recorremos para obtener la última matrícula registrada (aunque sea inactiva)
    for est in estudiantes_queryset.order_by('-FechaRetiro'):
        # Obtenemos la última matrícula que tuvo el estudiante
        ultima_matricula = est.matricula_set.order_by('-FechaMatricula').first()
        
        datos = {
            'estudiante': est,
            'matricula_activa': ultima_matricula, # Usamos el mismo nombre de variable para el template
        }
        lista_inactivos_completa.append(datos)

    # Paginación
    page_number = request.GET.get("page")
    paginator = Paginator(lista_inactivos_completa, 10)
    
    try:
        page_obj = paginator.page(page_number)
    except (PageNotAnInteger, EmptyPage):
        page_obj = paginator.page(1)
        
    # Rango de paginación (estilo profesional)
    paginator_range = range(max(1, page_obj.number - 2), min(paginator.num_pages, page_obj.number + 2) + 1)
    
    query_params = request.GET.copy()
    query_params.pop('page', None)
        
    context = {
        "estudiantes": page_obj, 
        "query": query,
        "paginator_range": paginator_range,
        "get_params": query_params.urlencode(),
        "total_inactivos": estudiantes_queryset.count()
    }
    return render(request, "matricula/listar_inactivos.html", context)

@login_required(login_url='/login/')
@permission_required ('Matricula.activar_estudiantes_nueva', raise_exception=True)
def reactivar_estudiante(request, pk):
    estudiante = get_object_or_404(Estudiantes_Nueva, EstudianteID=pk)
    estudiante.activo = True
    estudiante.FechaRetiro = None  # Limpiamos la fecha de retiro al reactivar
    estudiante.save()
    messages.success(request, f"El estudiante {estudiante.NombreCompleto} ha sido reactivado correctamente.")
    return redirect('Matricula:listar_estudiantes_inactivos')


@login_required(login_url='/login/')
@permission_required ('Matricula.ver_matricula', raise_exception=True)
def ver_hoja_matricula(request, matricula_id):
    matricula = get_object_or_404(Matricula, MatriculaID=matricula_id)
    estudiante = matricula.EstudianteID
    inscripciones = Inscripciones.objects.filter(MatriculaID=matricula)
    
    # Dividir el nombre completo en partes para el template
    nombre_completo_partes = estudiante.NombreCompleto.split(' ')
    primer_nombre = nombre_completo_partes[0] if len(nombre_completo_partes) > 0 else ''
    segundo_nombre = nombre_completo_partes[1] if len(nombre_completo_partes) > 1 else ''
    primer_apellido = nombre_completo_partes[2] if len(nombre_completo_partes) > 2 else ''
    segundo_apellido = nombre_completo_partes[3] if len(nombre_completo_partes) > 3 else ''

    context = {
        'matricula': matricula,
        'estudiante': estudiante,
        'inscripciones': inscripciones,
        'primer_nombre': primer_nombre,
        'segundo_nombre': segundo_nombre,
        'primer_apellido': primer_apellido,
        'segundo_apellido': segundo_apellido,
    }
    
    return render(request, 'Matricula/hoja_matricula.html', context)

@permission_required('Matricula.editar_estudiantes_nueva', raise_exception=True)
def editar_estudiante(request, estudiante_id):
    estudiante = get_object_or_404(Estudiantes_Nueva, EstudianteID=estudiante_id)
    
    if request.method == 'POST':
        form = EstudianteForm(request.POST, request.FILES, instance=estudiante)
        if form.is_valid():
            form.save()
            messages.success(request, 'El estudiante se ha actualizado correctamente.')
            return redirect('Matricula:listar_estudiantes')
        else:
            messages.error(request, 'Hubo un error al editar el estudiante. Por favor, revisa los datos.')
            # Si el formulario no es válido, se vuelve a renderizar la página con el formulario y los errores.
            print("Errores del formulario:", form.errors) # Agrega esta línea para ver los errores en la consola del servidor.
            
    else: # Petición GET
        form = EstudianteForm(instance=estudiante)
        
    return render(request, 'Matricula/editar_estudiante.html', {'form': form, 'estudiante': estudiante})




# --- VISTA DE FILTROS (PÁGINA 1) ---
# (Esta vista ahora solo necesita cargar los niveles)
@login_required(login_url='/login/')
@permission_required('Matricula.registro_notas', raise_exception=True)
def ingreso_calificaciones_filtro(request):
    """
    Muestra la interfaz para seleccionar los filtros (Nivel, Sección, Materia).
    """
    niveles = Nivel.objects.filter(Activo=True).order_by('NombreNivel')
    context = {
        'niveles': niveles,
    }
    return render(request, 'Matricula/ingreso_calificaciones_filtro.html', context)


# --- API PARA FILTROS (PÁGINA 1) ---
# (Esta vista carga Secciones Y Materias basado en Nivel)
@login_required(login_url='/login/')
@require_http_methods(["GET"])
def api_obtener_filtros_dinamicos(request):
    """
    API que devuelve Secciones y Materias basadas en el Nivel seleccionado.
    """
    nivel_id = request.GET.get('nivel_id')
    data = {'secciones': [], 'materias': []}
    
    if nivel_id:
        try:
            # 1. Obtener Secciones para el Nivel
            secciones = Seccion.objects.filter(NivelID=nivel_id, Activo=True).order_by('NombreSeccion')
            data['secciones'] = [{'id': s.SeccionID, 'nombre': s.NombreSeccion} for s in secciones]
            
            # 2. Obtener Materias para el Nivel
            materias_qs = Materias.objects.filter(NivelID=nivel_id, activo=True).order_by('NombreMateria')
            data['materias'] = [{'id': m.MateriaID, 'nombre': m.NombreMateria} for m in materias_qs]
            
        except Exception as e:
            print(f"Error en api_obtener_filtros_dinamicos: {e}")
            return JsonResponse({'error': 'Error interno del servidor'}, status=500)
            
    return JsonResponse(data, safe=False)


# --- ¡NUEVA VISTA API! ---
# --- API PARA CORTES (PÁGINA 2) ---
@login_required(login_url='/login/')
@require_http_methods(["GET"])
def api_get_cortes_por_semestre(request):
    """
    API que devuelve los Cortes basados en el Semestre seleccionado.
    """
    semestre_id = request.GET.get('semestre_id')
    data = []
    if semestre_id:
        try:
            # Filtra Cortes por SemestreID (basado en tu diagrama)
            cortes = Cortes.objects.filter(SemestreID=semestre_id, activo=True).order_by('NombreCorte')
            data = [{'id': c.CorteID, 'nombre': c.NombreCorte} for c in cortes]
        except Exception as e:
            print(f"Error en api_get_cortes_por_semestre: {e}")
            return JsonResponse({'error': 'Error interno'}, status=500)
    return JsonResponse(data, safe=False)


@login_required(login_url='/login/')
@require_http_methods(["GET", "POST"])
def ingreso_calificaciones(request, nivel, seccion, materia): # <-- 3 args
    """
    Muestra la lista de estudiantes y los filtros de Semestre/Corte.
    Permite ingresar/actualizar sus notas.
    """
    
    # 1. Recuperar objetos base (de la URL)
    try:
        nivel_obj = get_object_or_404(Nivel, pk=nivel)
        seccion_obj = get_object_or_404(Seccion, pk=seccion)
        materia_obj = get_object_or_404(Materias, pk=materia)
    except Exception:
        messages.error(request, 'Filtros de Nivel/Sección/Materia no válidos.')
        return redirect('Matricula:ingreso_calificaciones_filtro')

    # 2. Obtener Semestre y Corte (de POST o GET)
    semestre_id = None
    corte_id = None
    
    if request.method == 'POST':
        semestre_id = request.POST.get('semestre_id')
        corte_id = request.POST.get('corte_id')
    else: # GET
        semestre_id = request.GET.get('semestre_id')
        corte_id = request.GET.get('corte_id')

    # 3. Lógica de guardado (POST)
    if request.method == 'POST':
        if not semestre_id or not corte_id:
             messages.error(request, 'Se intentó guardar sin un Semestre o Corte válido.')
             return redirect(request.path) 
        
        try:
            semestre_obj = get_object_or_404(Semestres, pk=semestre_id)
            corte_obj = get_object_or_404(Cortes, pk=corte_id)
            anio_obj = timezone.now().year # Año automático
            
            estudiantes_activos = Matricula.objects.filter(
                NivelID=nivel_obj, SeccionID=seccion_obj, AñoMatricula=anio_obj, activo=True
            ).values_list('EstudianteID__EstudianteID', flat=True)
            
            inscripciones_qs = Inscripciones.objects.filter(
                EstudianteID__in=estudiantes_activos, MateriaID=materia_obj,
                NivelID=nivel_obj, SeccionID=seccion_obj, AñoMatricula=anio_obj
            )
            
            # (Pre-cargamos notas existentes para saber si es UPDATE O CREATE)
            calificaciones_actuales = {
                cal.InscripcionID.InscripcionID: cal
                for cal in Calificaciones.objects.filter(
                    InscripcionID__in=inscripciones_qs,
                    CorteID=corte_obj, SemestreID=semestre_obj,
                    MateriaID=materia_obj, # <--- CORREGIDO AQUÍ (1 de 3)
                    AnioMatricula=anio_obj # Usamos el nombre del modelo (con 'n')
                )
            }

            with transaction.atomic():
                for inscripcion in inscripciones_qs:
                    inscripcion_id = inscripcion.InscripcionID
                    nota_cuantitativa = request.POST.get(f'nota_{inscripcion_id}')
                    nota_cualitativa = request.POST.get(f'comentario_{inscripcion_id}')
                    
                    nota_float = None
                    if nota_cuantitativa:
                        try:
                            nota_float = float(nota_cuantitativa)
                            if not 0 <= nota_float <= 100: raise ValueError("Nota fuera de rango")
                        except (ValueError, TypeError):
                            messages.error(request, f'La nota para {inscripcion.EstudianteID.NombreCompleto} no es válida (0-100).')
                            params = f"?semestre_id={semestre_id}&corte_id={corte_id}"
                            return redirect(request.path + params)

                    calificacion = calificaciones_actuales.get(inscripcion_id)
                    
                    if calificacion: # Actualizar
                        calificacion.NotaCuantitativa = nota_float
                        calificacion.NotaCualitativa = nota_cualitativa
                        calificacion.save()
                    else: # Crear
                        if nota_float is not None or nota_cualitativa:
                            Calificaciones.objects.create(
                                InscripcionID=inscripcion, SemestreID=semestre_obj,
                                CorteID=corte_obj, 
                                MateriaID=materia_obj, # <--- CORREGIDO AQUÍ (2 de 3)
                                NotaCuantitativa=nota_float, 
                                NotaCualitativa=nota_cualitativa,
                                AnioMatricula=anio_obj, # Usamos el nombre del modelo (con 'n')
                                FechaRegistro=timezone.now(),
                            )
                            
            messages.success(request, 'Calificaciones guardadas exitosamente.')
            params = f"?semestre_id={semestre_id}&corte_id={corte_id}"
            return redirect(request.path + params)

        except Exception as e:
            messages.error(request, f'Ocurrió un error al guardar: {e}')
            return redirect(request.path)


    # 4. Lógica de Carga (GET)
    anio = timezone.now().year
    
    estudiantes_activos = Matricula.objects.filter(
        NivelID=nivel_obj, SeccionID=seccion_obj, AñoMatricula=anio, activo=True
    ).values_list('EstudianteID__EstudianteID', flat=True)

    inscripciones_qs = Inscripciones.objects.filter(
        EstudianteID__in=estudiantes_activos, MateriaID=materia_obj,
        NivelID=nivel_obj, SeccionID=seccion_obj, AñoMatricula=anio
    ).select_related('EstudianteID').order_by('EstudianteID__NombreCompleto')

    # 5. Preparar Contexto
    lista_alumnos = []
    calificaciones_actuales = {}
    semestre_obj = None
    corte_obj = None
    
    todos_semestres = Semestres.objects.filter(activo=True)
    todos_cortes = Cortes.objects.none() 

    if semestre_id and corte_id:
        try:
            semestre_obj = get_object_or_404(Semestres, pk=semestre_id)
            corte_obj = get_object_or_404(Cortes, pk=corte_id)
            todos_cortes = Cortes.objects.filter(SemestreID=semestre_obj, activo=True) 

            # (Cargamos las notas existentes SÓLO si hay semestre y corte)
            q_calificaciones = Calificaciones.objects.filter(
                InscripcionID__in=inscripciones_qs,
                CorteID=corte_obj, SemestreID=semestre_obj,
                MateriaID=materia_obj, # <--- CORREGIDO AQUÍ (3 de 3)
                AnioMatricula=anio    # Usamos el nombre del modelo (con 'n')
            )
            for cal in q_calificaciones:
                calificaciones_actuales[cal.InscripcionID.InscripcionID] = cal
        
        except Exception as e:
            # (El error que te salía ocurría en esta línea de arriba)
            messages.error(request, f"Error al cargar semestre o corte: {e}")

    # (Construimos la lista de alumnos para la tabla)
    for inscripcion in inscripciones_qs:
        inscripcion_id = inscripcion.InscripcionID
        calificacion = calificaciones_actuales.get(inscripcion_id) 
        
        lista_alumnos.append({
            'EstudianteID': inscripcion.EstudianteID.EstudianteID,
            'NombreCompleto': inscripcion.EstudianteID.NombreCompleto,
            'InscripcionID': inscripcion_id,
            'NotaCuantitativa': calificacion.NotaCuantitativa if calificacion and calificacion.NotaCuantitativa is not None else '',
            'NotaCualitativa': calificacion.NotaCualitativa if calificacion else '',
        })

    context = {
        'anio': anio,
        'nivel_obj': nivel_obj,
        'seccion_obj': seccion_obj,
        'materia_obj': materia_obj,
        'lista_alumnos': lista_alumnos,
        
        'todos_semestres': todos_semestres,
        'todos_cortes': todos_cortes,
        'selected_semestre_id': int(semestre_id) if semestre_id else None,
        'selected_corte_id': int(corte_id) if corte_id else None,
        'show_inputs': bool(semestre_id and corte_id) 
    }
    
    return render(request, 'Matricula/ingreso_calificaciones.html', context)

def get_nota_cualitativa(nota):
    """Convierte una nota cuantitativa (0-100) a su equivalente cualitativo."""
    if nota is None:
        return ""
    try:
        # Usamos Decimal para evitar problemas de precisión con float
        valor = Decimal(nota) 
        if valor >= 90: return "AA"
        if valor >= 76: return "AS"
        if valor >= 60: return "AF"
        if valor >= 0: return "AI" # Aseguramos que cubra de 0 a 59.99
        return "" # Para casos inesperados (negativos?)
    except:
        return "" # Si la conversión falla


# --- VISTA 1: FILTRO PARA EL REPORTE (Sin cambios) ---
@login_required(login_url='/login/')
@permission_required('Matricula.sabana_calificaciones', raise_exception=True) # Ajusta el permiso si es necesario
def reporte_calificaciones_filtro(request):
    """Muestra el formulario para seleccionar Nivel y Sección para el reporte."""
    niveles = Nivel.objects.filter(Activo=True).order_by('NombreNivel')
    # No necesitamos secciones aquí, se cargan con JS
    context = {
        'niveles': niveles,
    }
    return render(request, 'Matricula/reporte_calificaciones_filtro.html', context)


# --- VISTA 2: REPORTE VERTICAL (PARA LA PÁGINA WEB) ---
@login_required(login_url='/login/')
def ver_reporte_calificaciones(request, nivel_id, seccion_id):
    """
    Recopila y muestra el reporte de notas en formato VERTICAL (agrupado por materia)
    para la visualización en la página web.
    """
    try:
        nivel = get_object_or_404(Nivel, pk=nivel_id)
        seccion = get_object_or_404(Seccion, pk=seccion_id)
    except:
        messages.error(request, "El nivel o sección seleccionados no son válidos.")
        return redirect('Matricula:reporte_calificaciones_filtro')

    anio_actual = timezone.now().year

    try:
        c1 = Cortes.objects.get(NombreCorte__icontains='Primer Corte', activo=True)
        c2 = Cortes.objects.get(NombreCorte__icontains='Segundo Corte', activo=True)
        c3 = Cortes.objects.get(NombreCorte__icontains='Tercer Corte', activo=True)
        c4 = Cortes.objects.get(NombreCorte__icontains='Cuarto Corte', activo=True)
    except Cortes.DoesNotExist:
        messages.error(request, "No se encontraron los 4 cortes evaluativos necesarios (Primer, Segundo, Tercer, Cuarto). Por favor, créelos en Configuración.")
        return redirect('Matricula:reporte_calificaciones_filtro')

    estudiantes_activos = Estudiantes_Nueva.objects.filter(
        matricula__NivelID=nivel,
        matricula__SeccionID=seccion,
        matricula__AñoMatricula=anio_actual,
        matricula__activo=True
    ).order_by('NombreCompleto').distinct()

    materias_del_nivel = Materias.objects.filter(NivelID=nivel, activo=True).order_by('NombreMateria')

    calificaciones_map = {
        (cal.InscripcionID.EstudianteID_id, cal.InscripcionID.MateriaID_id, cal.CorteID_id): cal.NotaCuantitativa
        for cal in Calificaciones.objects.filter(
            InscripcionID__EstudianteID__in=estudiantes_activos,
            AnioMatricula=anio_actual
        ).select_related('InscripcionID') # Optimización
    }

    reporte_data = []
    for materia in materias_del_nivel:
        materia_info = {
            "materia": materia,
            "alumnos_data": []
        }
        for estudiante in estudiantes_activos:
            # Buscar notas
            n1 = calificaciones_map.get((estudiante.EstudianteID, materia.MateriaID, c1.CorteID))
            n2 = calificaciones_map.get((estudiante.EstudianteID, materia.MateriaID, c2.CorteID))
            n3 = calificaciones_map.get((estudiante.EstudianteID, materia.MateriaID, c3.CorteID))
            n4 = calificaciones_map.get((estudiante.EstudianteID, materia.MateriaID, c4.CorteID))

            # Calcular promedios usando Decimal
            s1, s2, nf = None, None, None
            n1_dec = Decimal(n1) if n1 is not None else None
            n2_dec = Decimal(n2) if n2 is not None else None
            n3_dec = Decimal(n3) if n3 is not None else None
            n4_dec = Decimal(n4) if n4 is not None else None

            if n1_dec is not None and n2_dec is not None:
                s1 = ((n1_dec + n2_dec) / 2).quantize(Decimal('0'), rounding=ROUND_HALF_UP) # Redondear a entero
            if n3_dec is not None and n4_dec is not None:
                s2 = ((n3_dec + n4_dec) / 2).quantize(Decimal('0'), rounding=ROUND_HALF_UP)
            if s1 is not None and s2 is not None:
                nf = ((s1 + s2) / 2).quantize(Decimal('0'), rounding=ROUND_HALF_UP)
            
            # Guardamos los Decimal (o None) para el template
            materia_info["alumnos_data"].append({
                "alumno": estudiante,
                "notas": {
                    "c1": {"cuant": n1_dec, "cual": get_nota_cualitativa(n1_dec)},
                    "c2": {"cuant": n2_dec, "cual": get_nota_cualitativa(n2_dec)},
                    "s1": {"cuant": s1, "cual": get_nota_cualitativa(s1)},
                    "c3": {"cuant": n3_dec, "cual": get_nota_cualitativa(n3_dec)},
                    "c4": {"cuant": n4_dec, "cual": get_nota_cualitativa(n4_dec)},
                    "s2": {"cuant": s2, "cual": get_nota_cualitativa(s2)},
                    "nf": {"cuant": nf, "cual": get_nota_cualitativa(nf)},
                }
            })
        # Incluimos la materia incluso si no tiene notas
        reporte_data.append(materia_info)

    context = {
        'nivel': nivel,
        'seccion': seccion,
        'reporte_data': reporte_data, # Datos para la vista vertical
        'anio': anio_actual,
    }
    return render(request, 'Matricula/ver_reporte_calificaciones.html', context)


# --- VISTA 3: EXPORTAR REPORTE HORIZONTAL A EXCEL ---
@login_required(login_url='/login/')
def exportar_reporte_excel(request, nivel_id, seccion_id):
    """
    Genera un archivo Excel (.xlsx) con el reporte de notas en formato HORIZONTAL.
    """
    try:
        nivel = get_object_or_404(Nivel, pk=nivel_id)
        seccion = get_object_or_404(Seccion, pk=seccion_id)
    except:
        messages.error(request, "El nivel o sección seleccionados no son válidos para exportar.")
        # Redirige de vuelta a la vista vertical si hay error
        return redirect('Matricula:ver_reporte_calificaciones', nivel_id=nivel_id, seccion_id=seccion_id) 

    anio_actual = timezone.now().year

    # Obtener los 4 cortes específicos por nombre (como antes)
    try:
        c1 = Cortes.objects.get(NombreCorte__icontains='Primer Corte', activo=True)
        c2 = Cortes.objects.get(NombreCorte__icontains='Segundo Corte', activo=True)
        c3 = Cortes.objects.get(NombreCorte__icontains='Tercer Corte', activo=True)
        c4 = Cortes.objects.get(NombreCorte__icontains='Cuarto Corte', activo=True)
    except Cortes.DoesNotExist:
        messages.error(request, "No se encontraron los 4 cortes evaluativos necesarios para exportar.")
        return redirect('Matricula:ver_reporte_calificaciones', nivel_id=nivel_id, seccion_id=seccion_id)

    # Obtener estudiantes y materias (como antes)
    estudiantes = Estudiantes_Nueva.objects.filter(
        matricula__NivelID=nivel,
        matricula__SeccionID=seccion,
        matricula__AñoMatricula=anio_actual,
        matricula__activo=True
    ).order_by('NombreCompleto').distinct()
    
    materias = Materias.objects.filter(NivelID=nivel, activo=True).order_by('NombreMateria')
    
    # Precargar calificaciones (como antes)
    calificaciones_map = {
        (cal.InscripcionID.EstudianteID_id, cal.InscripcionID.MateriaID_id, cal.CorteID_id): cal.NotaCuantitativa
        for cal in Calificaciones.objects.filter(
            InscripcionID__EstudianteID__in=estudiantes,
            AnioMatricula=anio_actual
        ).select_related('InscripcionID')
    }

    # --- INICIO DE LA GENERACIÓN DEL EXCEL ---
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    # Nombre de archivo dinámico
    filename = f"SabanaCalificaciones_{nivel.NombreNivel.replace(' ','_')}_{seccion.NombreSeccion}_{anio_actual}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{nivel.NombreNivel} - {seccion.NombreSeccion}"

    # --- Estilos ---
    font_bold = Font(bold=True, name='Calibri')
    font_header = Font(bold=True, size=12, name='Calibri')
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # --- Título ---
    ws.merge_cells('A1:C1')
    ws['A1'] = f"Boletín de Notas - {anio_actual}"
    ws['A1'].font = font_header
    ws.merge_cells('A2:C2')
    ws['A2'] = f"Nivel: {nivel.NombreNivel} - Sección: {seccion.NombreSeccion}"
    ws['A2'].font = font_header

    # --- Cabecera Compleja (3 filas) ---
    current_row = 4 # Empezamos en la fila 4
    
    # Fila 1 (Nombres de Materias)
    ws.cell(row=current_row, column=1, value="Estudiante").font = font_bold
    ws.cell(row=current_row, column=1).alignment = center_align
    ws.cell(row=current_row, column=1).border = thin_border
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row+2, end_column=1) # Unir verticalmente
    
    col_idx = 2
    for materia in materias:
        start_col = col_idx
        cell = ws.cell(row=current_row, column=start_col, value=materia.NombreMateria)
        cell.font = font_bold
        cell.alignment = center_align
        ws.merge_cells(start_row=current_row, start_column=start_col, end_row=current_row, end_column=start_col + 13)
        # Aplicar borde a la celda combinada
        for r in range(current_row, current_row + 1):
             for c_ in range(start_col, start_col + 14): ws.cell(row=r, column=c_).border = thin_border
        col_idx += 14
    current_row += 1

    # Fila 2 (Cortes y Semestres)
    col_idx = 2
    for materia in materias:
        headers2 = [("I CORTE", 2), ("II CORTE", 2), ("I SEM.", 2), ("III CORTE", 2), ("IV CORTE", 2), ("II SEM.", 2), ("FINAL", 2)]
        sub_col = col_idx
        for text, width in headers2:
            cell = ws.cell(row=current_row, column=sub_col, value=text)
            cell.font = font_bold
            cell.alignment = center_align
            ws.merge_cells(start_row=current_row, start_column=sub_col, end_row=current_row, end_column=sub_col + width - 1)
            # Aplicar borde
            for c_ in range(sub_col, sub_col + width): ws.cell(row=current_row, column=c_).border = thin_border
            sub_col += width
        col_idx += 14
    current_row += 1

    # Fila 3 (Nota y Equivalencia)
    col_idx = 2
    for materia in materias:
        for i in range(7):
            ws.cell(row=current_row, column=col_idx + (i*2), value="Nota").font = font_bold
            ws.cell(row=current_row, column=col_idx + (i*2)).alignment = center_align
            ws.cell(row=current_row, column=col_idx + (i*2)).border = thin_border
            ws.cell(row=current_row, column=col_idx + (i*2) + 1, value="Equiv.").font = font_bold
            ws.cell(row=current_row, column=col_idx + (i*2) + 1).alignment = center_align
            ws.cell(row=current_row, column=col_idx + (i*2) + 1).border = thin_border
        col_idx += 14
    current_row += 1

    # --- Llenar datos de los estudiantes ---
    for estudiante in estudiantes:
        ws.cell(row=current_row, column=1, value=estudiante.NombreCompleto).border = thin_border
        col_idx = 2
        for materia in materias:
            # Calcular notas (usando Decimal y redondeando)
            n1 = calificaciones_map.get((estudiante.EstudianteID, materia.MateriaID, c1.CorteID))
            n2 = calificaciones_map.get((estudiante.EstudianteID, materia.MateriaID, c2.CorteID))
            n3 = calificaciones_map.get((estudiante.EstudianteID, materia.MateriaID, c3.CorteID))
            n4 = calificaciones_map.get((estudiante.EstudianteID, materia.MateriaID, c4.CorteID))
            s1, s2, nf = None, None, None
            n1_dec = Decimal(n1) if n1 is not None else None
            n2_dec = Decimal(n2) if n2 is not None else None
            n3_dec = Decimal(n3) if n3 is not None else None
            n4_dec = Decimal(n4) if n4 is not None else None

            if n1_dec is not None and n2_dec is not None: s1 = ((n1_dec + n2_dec) / 2).quantize(Decimal('0'), rounding=ROUND_HALF_UP)
            if n3_dec is not None and n4_dec is not None: s2 = ((n3_dec + n4_dec) / 2).quantize(Decimal('0'), rounding=ROUND_HALF_UP)
            if s1 is not None and s2 is not None: nf = ((s1 + s2) / 2).quantize(Decimal('0'), rounding=ROUND_HALF_UP)
            
            # Formatear a entero para Excel (o dejar '-')
            n1_fmt = int(n1_dec) if n1_dec is not None else "-"
            n2_fmt = int(n2_dec) if n2_dec is not None else "-"
            s1_fmt = int(s1) if s1 is not None else "-"
            n3_fmt = int(n3_dec) if n3_dec is not None else "-"
            n4_fmt = int(n4_dec) if n4_dec is not None else "-"
            s2_fmt = int(s2) if s2 is not None else "-"
            nf_fmt = int(nf) if nf is not None else "-"

            # Poner datos en las celdas
            notas_fila_cuant = [n1_fmt, n2_fmt, s1_fmt, n3_fmt, n4_fmt, s2_fmt, nf_fmt]
            notas_fila_cual_orig = [n1_dec, n2_dec, s1, n3_dec, n4_dec, s2, nf] # Usar Decimal para get_nota_cualitativa

            for i, nota_cuant in enumerate(notas_fila_cuant):
                cell_nota = ws.cell(row=current_row, column=col_idx + (i*2), value=nota_cuant)
                cell_nota.alignment = center_align
                cell_nota.border = thin_border
                cell_cual = ws.cell(row=current_row, column=col_idx + (i*2) + 1, value=get_nota_cualitativa(notas_fila_cual_orig[i]))
                cell_cual.alignment = center_align
                cell_cual.border = thin_border
            col_idx += 14
        current_row += 1

    # --- Ajustar ancho de columnas ---
    ws.column_dimensions['A'].width = 35 # Ancho para nombres de estudiantes
    # Ancho estándar para notas y equivalencias
    for i in range(2, col_idx): 
        ws.column_dimensions[get_column_letter(i)].width = 7

    # --- Guardar y devolver ---
    wb.save(response)
    return response


def get_nota_cualitativa(nota):
    """Convierte una nota cuantitativa (0-100) a su equivalente cualitativo."""
    if nota is None:
        return ""
    try:
        valor = Decimal(nota) 
        if valor >= 90: return "AA"
        if valor >= 76: return "AS"
        if valor >= 60: return "AF"
        if valor >= 0: return "AI" 
        return "" 
    except:
        return ""


# --- VISTA 4: FILTRO PARA BOLETÍN INDIVIDUAL ---
@login_required(login_url='/login/')
def boletin_individual_filtro(request):
    """Muestra el formulario para seleccionar estudiante y periodos del boletín."""
    # No pasamos estudiantes aquí, se buscan con la API
    context = {}
    return render(request, 'Matricula/boletin_individual_filtro.html', context)


# --- VISTA 5: API PARA BUSCAR ESTUDIANTES (AUTOCOMPLETE) ---
@login_required(login_url='/login/')
def api_buscar_estudiantes(request):
    """Busca estudiantes por nombre (activos y con matrícula actual) y devuelve resultados para autocomplete."""
    query = request.GET.get('term', '')
    anio_actual = timezone.now().year
    results = []

    if len(query) >= 2: # Empezar a buscar con al menos 2 caracteres
        # Buscamos estudiantes activos Y que tengan matrícula activa en el año actual
        estudiantes = Estudiantes_Nueva.objects.filter(
            Q(NombreCompleto__icontains=query) &
            Q(activo=True) &
            Q(matricula__AñoMatricula=anio_actual) &
            Q(matricula__activo=True)
        ).distinct().order_by('NombreCompleto')[:10] # Limitar a 10 resultados

        for est in estudiantes:
            # Formato {label, value, id} para jQuery UI Autocomplete
            results.append({
                'id': est.EstudianteID, 
                'label': est.NombreCompleto, # Lo que ve el usuario
                'value': est.NombreCompleto  # Lo que se pone en el input al seleccionar
            })
            
    return JsonResponse(results, safe=False)


def _obtener_datos_boletin_context(estudiante, anio_actual, show_params, cortes_objs):
    """
    Función auxiliar para calcular todas las notas y promedios de un estudiante.
    Retorna un diccionario con 'boletin_data' y 'promedios'.
    """
    nivel = show_params['nivel']
    c1_obj, c2_obj, c3_obj, c4_obj = cortes_objs
    
    show_c1 = show_params.get('show_c1')
    show_c2 = show_params.get('show_c2')
    show_s1 = show_params.get('show_s1')
    show_c3 = show_params.get('show_c3')
    show_c4 = show_params.get('show_c4')
    show_s2 = show_params.get('show_s2')
    show_nf = show_params.get('show_nf')

    materias_inscritas = Materias.objects.filter(
        inscripciones__EstudianteID=estudiante,
        inscripciones__AñoMatricula=anio_actual,
        NivelID=nivel
    ).distinct().order_by('NombreMateria')

    calificaciones_map = {
        (cal.InscripcionID.MateriaID_id, cal.CorteID_id): cal.NotaCuantitativa
        for cal in Calificaciones.objects.filter(
            InscripcionID__EstudianteID=estudiante,
            AnioMatricula=anio_actual
        ).select_related('InscripcionID')
    }

    boletin_data = []
    sumas = {'c1': Decimal(0), 'c2': Decimal(0), 's1': Decimal(0), 'c3': Decimal(0), 'c4': Decimal(0), 's2': Decimal(0), 'nf': Decimal(0)}
    contadores = {'c1': 0, 'c2': 0, 's1': 0, 'c3': 0, 'c4': 0, 's2': 0, 'nf': 0}

    for materia in materias_inscritas:
        materia_notas = {}
        n1 = calificaciones_map.get((materia.MateriaID, c1_obj.CorteID)) if c1_obj else None
        n2 = calificaciones_map.get((materia.MateriaID, c2_obj.CorteID)) if c2_obj else None
        n3 = calificaciones_map.get((materia.MateriaID, c3_obj.CorteID)) if c3_obj else None
        n4 = calificaciones_map.get((materia.MateriaID, c4_obj.CorteID)) if c4_obj else None
        
        s1, s2, nf = None, None, None
        n1_dec, n2_dec, n3_dec, n4_dec = (Decimal(n) if n is not None else None for n in [n1, n2, n3, n4])

        if (show_s1 or show_nf) and n1_dec is not None and n2_dec is not None: 
            s1 = ((n1_dec + n2_dec) / 2).quantize(Decimal('0'), rounding=ROUND_HALF_UP)
        if (show_s2 or show_nf) and n3_dec is not None and n4_dec is not None: 
            s2 = ((n3_dec + n4_dec) / 2).quantize(Decimal('0'), rounding=ROUND_HALF_UP)
        if show_nf and s1 is not None and s2 is not None: 
            nf = ((s1 + s2) / 2).quantize(Decimal('0'), rounding=ROUND_HALF_UP)

        if n1_dec is not None: sumas['c1'] += n1_dec; contadores['c1'] += 1
        if n2_dec is not None: sumas['c2'] += n2_dec; contadores['c2'] += 1
        if s1 is not None: sumas['s1'] += s1; contadores['s1'] += 1
        if n3_dec is not None: sumas['c3'] += n3_dec; contadores['c3'] += 1
        if n4_dec is not None: sumas['c4'] += n4_dec; contadores['c4'] += 1
        if s2 is not None: sumas['s2'] += s2; contadores['s2'] += 1
        if nf is not None: sumas['nf'] += nf; contadores['nf'] += 1

        if show_c1: materia_notas['c1'] = {"cuant": n1_dec, "cual": get_nota_cualitativa(n1_dec)}
        if show_c2: materia_notas['c2'] = {"cuant": n2_dec, "cual": get_nota_cualitativa(n2_dec)}
        if show_s1: materia_notas['s1'] = {"cuant": s1, "cual": get_nota_cualitativa(s1)}
        if show_c3: materia_notas['c3'] = {"cuant": n3_dec, "cual": get_nota_cualitativa(n3_dec)}
        if show_c4: materia_notas['c4'] = {"cuant": n4_dec, "cual": get_nota_cualitativa(n4_dec)}
        if show_s2: materia_notas['s2'] = {"cuant": s2, "cual": get_nota_cualitativa(s2)}
        if show_nf: materia_notas['nf'] = {"cuant": nf, "cual": get_nota_cualitativa(nf)}

        boletin_data.append({"materia": materia, "notas": materia_notas})

    promedios = {}
    for key in sumas:
        if contadores[key] > 0:
            promedio_dec = (sumas[key] / contadores[key]).quantize(Decimal('0'), rounding=ROUND_HALF_UP)
            promedios[key] = {"cuant": promedio_dec, "cual": get_nota_cualitativa(promedio_dec)}
        else:
            promedios[key] = {"cuant": None, "cual": "-"}
            
    return {
        'boletin_data': boletin_data,
        'promedios': promedios,
        'materias_count': materias_inscritas.count(),
        'notas_count': len(calificaciones_map)
    }


# --- VISTA 6: GENERAR Y MOSTRAR BOLETÍN INDIVIDUAL ---
@login_required(login_url='/login/')
def ver_boletin_individual(request, estudiante_id):
    """
    Recopila, calcula y muestra el boletín de notas individual
    según los periodos seleccionados en el filtro (parámetros GET).
    """
    try:
        estudiante = get_object_or_404(Estudiantes_Nueva, pk=estudiante_id)
    except:
        messages.error(request, "El estudiante seleccionado no es válido.")
        return redirect('Matricula:boletin_individual_filtro')

    anio_actual = timezone.now().year
    matricula_activa = Matricula.objects.filter(
        EstudianteID=estudiante, AñoMatricula=anio_actual, activo=True
    ).select_related('NivelID', 'SeccionID').first()

    if not matricula_activa:
        messages.warning(request, f"El estudiante {estudiante.NombreCompleto} no tiene una matrícula activa para el año {anio_actual}.")
        return redirect('Matricula:boletin_individual_filtro')

    nivel = matricula_activa.NivelID
    seccion = matricula_activa.SeccionID

    tipo_nota = request.GET.get('tipo_nota', 'ambas')
    show_cuant = tipo_nota in ['ambas', 'cuantitativa']
    show_cual = tipo_nota in ['ambas', 'cualitativa']
    colspan_periodo = 2 if (show_cuant and show_cual) else 1

    show_params = {
        'nivel': nivel,
        'show_c1': request.GET.get('c1') == 'on',
        'show_c2': request.GET.get('c2') == 'on',
        'show_s1': request.GET.get('s1') == 'on',
        'show_c3': request.GET.get('c3') == 'on',
        'show_c4': request.GET.get('c4') == 'on',
        'show_s2': request.GET.get('s2') == 'on',
        'show_nf': request.GET.get('nf') == 'on',
    }
    
    if request.GET.get('all') == 'on':
        for k in ['show_c1', 'show_c2', 'show_s1', 'show_c3', 'show_c4', 'show_s2', 'show_nf']:
            show_params[k] = True
    else:
        if show_params['show_s1']: show_params['show_c1'] = show_params['show_c2'] = True
        if show_params['show_s2']: show_params['show_c3'] = show_params['show_c4'] = True
        if show_params['show_nf']: show_params['show_c1'] = show_params['show_c2'] = show_params['show_s1'] = \
                                   show_params['show_c3'] = show_params['show_c4'] = show_params['show_s2'] = True

    try:
        c1 = Cortes.objects.get(NombreCorte__icontains='Primer Corte', activo=True)
        c2 = Cortes.objects.get(NombreCorte__icontains='Segundo Corte', activo=True)
        c3 = Cortes.objects.get(NombreCorte__icontains='Tercer Corte', activo=True)
        c4 = Cortes.objects.get(NombreCorte__icontains='Cuarto Corte', activo=True)
        cortes_objs = (c1, c2, c3, c4)
    except Cortes.DoesNotExist:
        messages.error(request, "Faltan uno o más cortes evaluativos necesarios en la configuración.")
        return redirect('Matricula:boletin_individual_filtro')

    data = _obtener_datos_boletin_context(estudiante, anio_actual, show_params, cortes_objs)

    context = {
        'estudiante': estudiante, 'nivel': nivel, 'seccion': seccion, 'anio': anio_actual,
        'boletin_data': data['boletin_data'],
        'promedios': data['promedios'],
        'show_c1': show_params['show_c1'], 'show_c2': show_params['show_c2'], 'show_s1': show_params['show_s1'],
        'show_c3': show_params['show_c3'], 'show_c4': show_params['show_c4'], 'show_s2': show_params['show_s2'], 'show_nf': show_params['show_nf'],
        'show_cuant': show_cuant, 'show_cual': show_cual, 'colspan_periodo': colspan_periodo,
    }
    return render(request, 'Matricula/boletin_individual.html', context)


# --- VISTA 7: IMPRIMIR BOLETÍN INDIVIDUAL (CON PROMEDIOS) ---
@login_required(login_url='/login/')
@login_required(login_url='/login/')
def imprimir_boletin_individual(request, estudiante_id):
    """
    Prepara los datos para el boletín individual imprimible.
    """
    try:
        estudiante = get_object_or_404(Estudiantes_Nueva, pk=estudiante_id)
    except:
        messages.error(request, "El estudiante seleccionado no es válido.")
        return redirect('Matricula:boletin_individual_filtro')

    anio_actual = timezone.now().year
    matricula_activa = Matricula.objects.filter(
        EstudianteID=estudiante, AñoMatricula=anio_actual, activo=True
    ).select_related('NivelID', 'SeccionID').first()

    if not matricula_activa:
        messages.warning(request, f"El estudiante {estudiante.NombreCompleto} no tiene una matrícula activa para el año {anio_actual}.")
        return redirect('Matricula:boletin_individual_filtro')

    nivel = matricula_activa.NivelID
    seccion = matricula_activa.SeccionID

    tipo_nota = request.GET.get('tipo_nota', 'ambas')
    show_cuant = tipo_nota in ['ambas', 'cuantitativa']
    show_cual = tipo_nota in ['ambas', 'cualitativa']
    colspan_periodo = 2 if (show_cuant and show_cual) else 1

    show_params = {
        'nivel': nivel,
        'show_c1': request.GET.get('c1') == 'on',
        'show_c2': request.GET.get('c2') == 'on',
        'show_s1': request.GET.get('s1') == 'on',
        'show_c3': request.GET.get('c3') == 'on',
        'show_c4': request.GET.get('c4') == 'on',
        'show_s2': request.GET.get('s2') == 'on',
        'show_nf': request.GET.get('nf') == 'on',
    }

    if request.GET.get('all') == 'on':
        for k in ['show_c1', 'show_c2', 'show_s1', 'show_c3', 'show_c4', 'show_s2', 'show_nf']:
            show_params[k] = True
    else:
        if show_params['show_s1']: show_params['show_c1'] = show_params['show_c2'] = True
        if show_params['show_s2']: show_params['show_c3'] = show_params['show_c4'] = True
        if show_params['show_nf']: show_params['show_c1'] = show_params['show_c2'] = show_params['show_s1'] = \
                                   show_params['show_c3'] = show_params['show_c4'] = show_params['show_s2'] = True

    try:
        c1 = Cortes.objects.get(NombreCorte__icontains='Primer Corte', activo=True)
        c2 = Cortes.objects.get(NombreCorte__icontains='Segundo Corte', activo=True)
        c3 = Cortes.objects.get(NombreCorte__icontains='Tercer Corte', activo=True)
        c4 = Cortes.objects.get(NombreCorte__icontains='Cuarto Corte', activo=True)
        cortes_objs = (c1, c2, c3, c4)
    except Cortes.DoesNotExist:
        messages.error(request, "Faltan uno o más cortes evaluativos necesarios en la configuración.")
        return redirect('Matricula:boletin_individual_filtro')

    data = _obtener_datos_boletin_context(estudiante, anio_actual, show_params, cortes_objs)

    context = {
        'estudiante': estudiante, 'nivel': nivel, 'seccion': seccion, 'anio': anio_actual,
        'boletin_data': data['boletin_data'],
        'promedios': data['promedios'],
        'show_c1': show_params['show_c1'], 'show_c2': show_params['show_c2'], 'show_s1': show_params['show_s1'],
        'show_c3': show_params['show_c3'], 'show_c4': show_params['show_c4'], 'show_s2': show_params['show_s2'], 'show_nf': show_params['show_nf'],
        'fecha_emision': timezone.now(),
        'show_cuant': show_cuant, 'show_cual': show_cual, 'colspan_periodo': colspan_periodo,
    }
    return render(request, 'Matricula/boletin_individual_print.html', context)


# --- VISTA 8: FILTRO PARA BOLETÍN GRUPAL (POR SECCIÓN) ---
@login_required(login_url='/login/')
def boletin_grupal_filtro(request):
    """
    Permite filtrar por grado y sección para obtener una lista de alumnos 
    y seleccionar para impresión masiva.
    """
    niveles = Nivel.objects.filter(Activo=True).order_by('NombreNivel')
    anio_actual = timezone.now().year
    
    # Parámetros de filtrado
    nivel_id = request.GET.get('nivel')
    seccion_id = request.GET.get('seccion')
    estudiantes_con_estado = []

    if nivel_id and seccion_id:
        try:
            # Obtener estudiantes matriculados en esa sección este año
            matriculados = Matricula.objects.filter(
                NivelID_id=nivel_id,
                SeccionID_id=seccion_id,
                AñoMatricula=anio_actual,
                activo=True
            ).select_related('EstudianteID')

            # Pre-cargar cortes para el chequeo de "completitud"
            show_c1 = request.GET.get('c1') == 'on'
            show_c2 = request.GET.get('c2') == 'on'
            show_c3 = request.GET.get('c3') == 'on'
            show_c4 = request.GET.get('c4') == 'on'
            show_all = request.GET.get('all') == 'on'
            
            # Calculamos cuántos periodos se están pidiendo para el chequeo de "completitud"
            cortes_activos_count = 0
            if show_all: 
                cortes_activos_count = 4
            else:
                if request.GET.get('c1') == 'on' or request.GET.get('s1') == 'on' or request.GET.get('nf') == 'on': cortes_activos_count += 1
                if request.GET.get('c2') == 'on' or request.GET.get('s1') == 'on' or request.GET.get('nf') == 'on': cortes_activos_count += 1
                if request.GET.get('c3') == 'on' or request.GET.get('s2') == 'on' or request.GET.get('nf') == 'on': cortes_activos_count += 1
                if request.GET.get('c4') == 'on' or request.GET.get('s2') == 'on' or request.GET.get('nf') == 'on': cortes_activos_count += 1

            # Si no hay periodos seleccionados, por defecto comparamos contra 1 corte para mostrar el total de materias
            if cortes_activos_count == 0:
                cortes_activos_count = 1

            n_materias = Materias.objects.filter(NivelID_id=nivel_id, activo=True).count()
            # n_esperadas es materias * periodos seleccionados
            n_esperadas = n_materias * cortes_activos_count

            for mat in matriculados:
                est = mat.EstudianteID
                # Chequeo de completitud
                n_reales = Calificaciones.objects.filter(InscripcionID__EstudianteID=est, AnioMatricula=anio_actual).count()
                
                estudiantes_con_estado.append({
                    'est': est,
                    'completo': n_reales >= n_esperadas if n_esperadas > 0 else True,
                    'reales': n_reales,
                    'esperadas': n_esperadas
                })
        except Exception as e:
            messages.error(request, f"Error al cargar estudiantes: {e}")

    context = {
        'niveles': niveles,
        'estudiantes_lista': estudiantes_con_estado,
        'nivel_id': int(nivel_id) if nivel_id else None,
        'seccion_id': int(seccion_id) if seccion_id else None,
        'tipo_nota_selected': request.GET.get('tipo_nota', 'ambas'),
    }
    return render(request, 'Matricula/boletin_grupal_filtro.html', context)


@login_required(login_url='/login/')
def imprimir_boletin_grupal(request):
    """
    Recibe una lista de IDs de estudiantes y genera la vista de impresión masiva.
    """
    if request.method == 'POST':
        estudiante_ids = request.POST.getlist('estudiante_ids')
    else:
        # Por si acaso alguien entra por GET, aunque se espera POST
        return redirect('Matricula:boletin_grupal_filtro')

    if not estudiante_ids:
        messages.warning(request, "No se seleccionaron estudiantes para imprimir.")
        return redirect('Matricula:boletin_grupal_filtro')

    anio_actual = timezone.now().year
    
    # Parámetros de visualización (igual que el individual)
    tipo_nota = request.POST.get('tipo_nota', 'ambas')
    show_cuant = tipo_nota in ['ambas', 'cuantitativa']
    show_cual = tipo_nota in ['ambas', 'cualitativa']
    colspan_periodo = 2 if (show_cuant and show_cual) else 1

    show_params_base = {
        'show_c1': request.POST.get('c1') == 'on',
        'show_c2': request.POST.get('c2') == 'on',
        'show_s1': request.POST.get('s1') == 'on',
        'show_c3': request.POST.get('c3') == 'on',
        'show_c4': request.POST.get('c4') == 'on',
        'show_s2': request.POST.get('s2') == 'on',
        'show_nf': request.POST.get('nf') == 'on',
    }

    if request.POST.get('all') == 'on':
        for k in ['show_c1', 'show_c2', 'show_s1', 'show_c3', 'show_c4', 'show_s2', 'show_nf']:
            show_params_base[k] = True
    else:
        if show_params_base['show_s1']: show_params_base['show_c1'] = show_params_base['show_c2'] = True
        if show_params_base['show_s2']: show_params_base['show_c3'] = show_params_base['show_c4'] = True
        if show_params_base['show_nf']: show_params_base['show_c1'] = show_params_base['show_c2'] = show_params_base['show_s1'] = \
                                        show_params_base['show_c3'] = show_params_base['show_c4'] = show_params_base['show_s2'] = True

    try:
        c1 = Cortes.objects.get(NombreCorte__icontains='Primer Corte', activo=True)
        c2 = Cortes.objects.get(NombreCorte__icontains='Segundo Corte', activo=True)
        c3 = Cortes.objects.get(NombreCorte__icontains='Tercer Corte', activo=True)
        c4 = Cortes.objects.get(NombreCorte__icontains='Cuarto Corte', activo=True)
        cortes_objs = (c1, c2, c3, c4)
    except Cortes.DoesNotExist:
        messages.error(request, "Configuración de Cortes incompleta.")
        return redirect('Matricula:boletin_grupal_filtro')

    lista_boletines = []
    
    # Procesar cada estudiante seleccionado
    for eid in estudiante_ids:
        try:
            estudiante = Estudiantes_Nueva.objects.get(pk=eid)
            matricula = Matricula.objects.filter(EstudianteID=estudiante, AñoMatricula=anio_actual, activo=True).first()
            
            if matricula:
                st_params = show_params_base.copy()
                st_params['nivel'] = matricula.NivelID
                
                res = _obtener_datos_boletin_context(estudiante, anio_actual, st_params, cortes_objs)
                
                lista_boletines.append({
                    'estudiante': estudiante,
                    'nivel': matricula.NivelID,
                    'seccion': matricula.SeccionID,
                    'boletin_data': res['boletin_data'],
                    'promedios': res['promedios']
                })
        except Estudiantes_Nueva.DoesNotExist:
            continue

    context = {
        'lista_boletines': lista_boletines,
        'anio': anio_actual,
        'fecha_emision': timezone.now(),
        'show_c1': show_params_base['show_c1'], 'show_c2': show_params_base['show_c2'], 'show_s1': show_params_base['show_s1'],
        'show_c3': show_params_base['show_c3'], 'show_c4': show_params_base['show_c4'], 'show_s2': show_params_base['show_s2'], 'show_nf': show_params_base['show_nf'],
        'show_cuant': show_cuant, 'show_cual': show_cual, 'colspan_periodo': colspan_periodo,
    }
    
    return render(request, 'Matricula/boletin_grupal_print.html', context)


# ---===================================================---
# ---   NUEVAS VISTAS PARA EL HISTORIAL ACADÉMICO      ---
# ---===================================================---

@login_required(login_url='/login/')
@permission_required('Matricula.historial_academico', raise_exception=True)
def historial_academico_filtro(request):
    """Muestra el formulario para buscar al estudiante para ver su historial."""
    context = {}
    return render(request, 'Matricula/historial_academico_filtro.html', context)

@login_required(login_url='/login/')
def ver_historial_academico(request, estudiante_id):
    """
    Recopila, calcula y muestra el historial completo de notas de un estudiante,
    agrupado por año académico (matrícula).
    """
    try:
        estudiante = get_object_or_404(Estudiantes_Nueva, pk=estudiante_id)
    except:
        messages.error(request, "El estudiante seleccionado no es válido.")
        return redirect('Matricula:historial_academico_filtro')

    matriculas_estudiante = Matricula.objects.filter(
        EstudianteID=estudiante
    ).select_related('NivelID', 'SeccionID').order_by('AñoMatricula')

    if not matriculas_estudiante.exists():
        messages.warning(request, f"El estudiante {estudiante.NombreCompleto} no tiene matrículas registradas.")
        return redirect('Matricula:historial_academico_filtro')

    try:
        c1_obj = Cortes.objects.get(NombreCorte__icontains='Primer Corte', activo=True)
        c2_obj = Cortes.objects.get(NombreCorte__icontains='Segundo Corte', activo=True)
        c3_obj = Cortes.objects.get(NombreCorte__icontains='Tercer Corte', activo=True)
        c4_obj = Cortes.objects.get(NombreCorte__icontains='Cuarto Corte', activo=True)
    except Cortes.DoesNotExist:
        messages.error(request, "Faltan cortes evaluativos en la configuración.")
        return redirect('Matricula:historial_academico_filtro')

    calificaciones_map = {
        (cal.AnioMatricula, cal.InscripcionID.MateriaID_id, cal.CorteID_id): cal.NotaCuantitativa
        for cal in Calificaciones.objects.filter(
            InscripcionID__EstudianteID=estudiante
        ).select_related('InscripcionID')
    }

    historial_por_anio = []
    for matricula in matriculas_estudiante:
        anio_lectivo = matricula.AñoMatricula
        nivel_anio = matricula.NivelID
        seccion_anio = matricula.SeccionID

        materias_inscritas_anio = Materias.objects.filter(
            inscripciones__EstudianteID=estudiante,
            inscripciones__AñoMatricula=anio_lectivo,
            NivelID=nivel_anio 
        ).distinct().order_by('NombreMateria')

        datos_anio_actual = {
            "anio": anio_lectivo, "nivel": nivel_anio, "seccion": seccion_anio, "boletin_data": []
        }

        for materia in materias_inscritas_anio:
            n1 = calificaciones_map.get((anio_lectivo, materia.MateriaID, c1_obj.CorteID))
            n2 = calificaciones_map.get((anio_lectivo, materia.MateriaID, c2_obj.CorteID))
            n3 = calificaciones_map.get((anio_lectivo, materia.MateriaID, c3_obj.CorteID))
            n4 = calificaciones_map.get((anio_lectivo, materia.MateriaID, c4_obj.CorteID))
            s1, s2, nf = None, None, None
            n1_dec, n2_dec, n3_dec, n4_dec = (Decimal(n) if n is not None else None for n in [n1, n2, n3, n4])

            if n1_dec is not None and n2_dec is not None: s1 = ((n1_dec + n2_dec) / 2).quantize(Decimal('0'), rounding=ROUND_HALF_UP)
            if n3_dec is not None and n4_dec is not None: s2 = ((n3_dec + n4_dec) / 2).quantize(Decimal('0'), rounding=ROUND_HALF_UP)
            if s1 is not None and s2 is not None: nf = ((s1 + s2) / 2).quantize(Decimal('0'), rounding=ROUND_HALF_UP)

            datos_anio_actual["boletin_data"].append({
                "materia": materia,
                "notas": {
                    "c1": {"cuant": n1_dec, "cual": get_nota_cualitativa(n1_dec)},
                    "c2": {"cuant": n2_dec, "cual": get_nota_cualitativa(n2_dec)},
                    "s1": {"cuant": s1, "cual": get_nota_cualitativa(s1)},
                    "c3": {"cuant": n3_dec, "cual": get_nota_cualitativa(n3_dec)},
                    "c4": {"cuant": n4_dec, "cual": get_nota_cualitativa(n4_dec)},
                    "s2": {"cuant": s2, "cual": get_nota_cualitativa(s2)},
                    "nf": {"cuant": nf, "cual": get_nota_cualitativa(nf)},
                }
            })
        historial_por_anio.append(datos_anio_actual)

    context = {
        'estudiante': estudiante,
        'historial_por_anio': historial_por_anio,
        'fecha_emision': timezone.now(),
    }
    return render(request, 'Matricula/historial_academico.html', context)

@login_required(login_url='/login/')
def imprimir_historial_anio(request, estudiante_id, anio, nivel_id, seccion_id):
    try:
        estudiante = get_object_or_404(Estudiantes_Nueva, pk=estudiante_id)
        matricula_anio = get_object_or_404(Matricula,
            EstudianteID=estudiante, AñoMatricula=anio,
            NivelID=nivel_id, SeccionID=seccion_id
        )
        nivel = matricula_anio.NivelID
        seccion = matricula_anio.SeccionID
    except Matricula.DoesNotExist:
         messages.warning(request, f"No se encontró la matrícula específica...")
         return redirect('Matricula:ver_historial_academico', estudiante_id=estudiante_id)
    except Exception as e:
        messages.error(request, f"Error al buscar datos: {e}")
        return redirect('Matricula:historial_academico_filtro')

    anio_lectivo = anio 
    try:
        c1_obj = Cortes.objects.get(NombreCorte__icontains='Primer Corte', activo=True)
        c2_obj = Cortes.objects.get(NombreCorte__icontains='Segundo Corte', activo=True)
        c3_obj = Cortes.objects.get(NombreCorte__icontains='Tercer Corte', activo=True)
        c4_obj = Cortes.objects.get(NombreCorte__icontains='Cuarto Corte', activo=True)
    except Cortes.DoesNotExist:
        messages.error(request, "Faltan cortes evaluativos.")
        return redirect('Matricula:ver_historial_academico', estudiante_id=estudiante_id)

    materias_inscritas_anio = Materias.objects.filter(
        inscripciones__EstudianteID=estudiante, inscripciones__AñoMatricula=anio_lectivo, 
        NivelID=nivel 
    ).distinct().order_by('NombreMateria')
    calificaciones_map = {
        (cal.InscripcionID.MateriaID_id, cal.CorteID_id): cal.NotaCuantitativa
        for cal in Calificaciones.objects.filter(
            InscripcionID__EstudianteID=estudiante, AnioMatricula=anio_lectivo 
        ).select_related('InscripcionID')
    }

    boletin_data_anio = []
    for materia in materias_inscritas_anio:
        n1 = calificaciones_map.get((materia.MateriaID, c1_obj.CorteID))
        n2 = calificaciones_map.get((materia.MateriaID, c2_obj.CorteID))
        n3 = calificaciones_map.get((materia.MateriaID, c3_obj.CorteID))
        n4 = calificaciones_map.get((materia.MateriaID, c4_obj.CorteID))
        s1, s2, nf = None, None, None
        n1_dec, n2_dec, n3_dec, n4_dec = (Decimal(n) if n is not None else None for n in [n1, n2, n3, n4])
        if n1_dec is not None and n2_dec is not None: s1 = ((n1_dec + n2_dec) / 2).quantize(Decimal('0'), rounding=ROUND_HALF_UP)
        if n3_dec is not None and n4_dec is not None: s2 = ((n3_dec + n4_dec) / 2).quantize(Decimal('0'), rounding=ROUND_HALF_UP)
        if s1 is not None and s2 is not None: nf = ((s1 + s2) / 2).quantize(Decimal('0'), rounding=ROUND_HALF_UP)

        boletin_data_anio.append({
            "materia": materia,
            "notas": {
                "c1": {"cuant": n1_dec, "cual": get_nota_cualitativa(n1_dec)},
                "c2": {"cuant": n2_dec, "cual": get_nota_cualitativa(n2_dec)},
                "s1": {"cuant": s1, "cual": get_nota_cualitativa(s1)},
                "c3": {"cuant": n3_dec, "cual": get_nota_cualitativa(n3_dec)},
                "c4": {"cuant": n4_dec, "cual": get_nota_cualitativa(n4_dec)},
                "s2": {"cuant": s2, "cual": get_nota_cualitativa(s2)},
                "nf": {"cuant": nf, "cual": get_nota_cualitativa(nf)},
            }
        })

    context = {
        'estudiante': estudiante, 'nivel': nivel, 'seccion': seccion, 'anio': anio_lectivo, 
        'boletin_data': boletin_data_anio, 
        'show_c1': True, 'show_c2': True, 'show_s1': True, # Asumimos mostrar todo
        'show_c3': True, 'show_c4': True, 'show_s2': True, 'show_nf': True,
        'fecha_emision': timezone.now(),
    }
    return render(request, 'Matricula/historial_academico_print.html', context)

