import csv
from django.shortcuts import render, get_object_or_404
from django.http import HttpResponse, JsonResponse 
from Matricula.models import Matricula, Inscripciones
from configuracion.models import Nivel, Seccion, Turno, Modalidades 
from django.core.paginator import Paginator
from django.views.decorators.http import require_GET
import openpyxl
from openpyxl.utils import get_column_letter
from datetime import datetime

# --- VISTA PRINCIPAL (sin cambios) ---
def reportes_home(request):
    """
    Vista principal para el dashboard de reportes.
    """
    context = {
        'titulo': 'Dashboard de Reportes'
    }
    return render(request, 'reportes/reportes_home.html', context)

def reporte_estudiantes_view(request):
    """
    Vista de reporte que solo muestra estudiantes matriculados Y activos.
    """
    # 1. FILTRO CORREGIDO: 
    # Aseguramos que la matrícula esté activa Y el estudiante NO esté dado de baja
    matriculas = Matricula.objects.filter(
        activo=True, 
        EstudianteID__activo=True  # <--- Este filtro evita que salgan los retirados
    ).select_related(
        'EstudianteID', 'NivelID', 'SeccionID', 'TurnoID', 'ModalidadID'
    )

    # 2. Obtener datos de filtros desde el GET
    selected_nivel = request.GET.get('nivel')
    selected_seccion = request.GET.get('seccion')
    selected_turno = request.GET.get('turno')
    selected_modalidad = request.GET.get('modalidad')

    # 3. Aplicar Filtros a la QuerySet si existen
    if selected_nivel:
        matriculas = matriculas.filter(NivelID__NivelID=selected_nivel)
    if selected_seccion:
        matriculas = matriculas.filter(SeccionID__SeccionID=selected_seccion)
    if selected_turno:
        matriculas = matriculas.filter(TurnoID__TurnoID=selected_turno)
    if selected_modalidad:
        matriculas = matriculas.filter(ModalidadID__ModalidadID=selected_modalidad)

    # --- LÓGICA DE CARDS (Contadores automáticos sobre el filtro anterior) ---
    total_estudiantes = matriculas.count()
    total_hombres = matriculas.filter(EstudianteID__Genero='M').count()
    total_mujeres = matriculas.filter(EstudianteID__Genero='F').count()

    # 4. Paginación (10 registros por página)
    # Ordenamos por nombre para que el reporte sea legible
    paginator = Paginator(matriculas.order_by('EstudianteID__NombreCompleto'), 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    # Obtener secciones si hay un nivel seleccionado (para mantener el filtro visual)
    secciones = []
    if selected_nivel:
        secciones = Seccion.objects.filter(NivelID__NivelID=selected_nivel)

    # 5. Contexto para el template
    context = {
        'titulo': 'Reporte General de Estudiantes Activos',
        'page_obj': page_obj,
        'total_estudiantes': total_estudiantes,
        'total_hombres': total_hombres,
        'total_mujeres': total_mujeres,
        'niveles': Nivel.objects.filter(Activo=True),
        'secciones': secciones, # Pasamos las secciones filtradas
        'turnos': Turno.objects.all(),
        'modalidades': Modalidades.objects.all(),
        'selected_nivel': selected_nivel,
        'selected_seccion': selected_seccion,
        'selected_turno': selected_turno,
        'selected_modalidad': selected_modalidad,
    }
    return render(request, 'reportes/reporte_estudiantes.html', context)

@require_GET
def get_secciones_por_nivel(request):
    nivel_id = request.GET.get('nivel_id')
    secciones_data = []
    if nivel_id:
        secciones_filtradas = Seccion.objects.filter(NivelID__NivelID=nivel_id).order_by('NombreSeccion')
        secciones_data = [{'id': s.SeccionID, 'nombre': s.NombreSeccion} for s in secciones_filtradas]
    return JsonResponse({'secciones': secciones_data})


# --- VISTAS RESTANTES (sin cambios) ---

def imprimir_reporte_estudiante_view(request, estudiante_id):
    """
    Vista para imprimir los detalles completos de un estudiante.
    """
    matricula = get_object_or_404(Matricula, EstudianteID__EstudianteID=estudiante_id, activo=True)
    inscripciones = Inscripciones.objects.filter(MatriculaID=matricula.MatriculaID).select_related('MateriaID', 'NivelID', 'SeccionID', 'ModalidadID', 'TurnoID')
    
    context = {
        'titulo': f'Reporte de {matricula.EstudianteID.NombreCompleto}',
        'matricula': matricula,
        'inscripciones': inscripciones,
    }
    
    return render(request, 'reportes/imprimir_reporte.html', context)


def imprimir_reporte_completo_view(request):
    """
    Vista optimizada para imprimir estudiantes filtrando automáticamente 
    por el año actual (2026) y criterios de búsqueda.
    """
    # 1. Obtener el año actual dinámicamente
    anio_actual = datetime.now().year
    
    # 2. Consulta base: Filtramos por activo=True y por el año actual
    # Esto evita que se mezclen datos de 2025 en el reporte de 2026
    matriculas = Matricula.objects.filter(
        activo=True, 
        AñoMatricula=anio_actual
    ).select_related(
        'EstudianteID', 'NivelID', 'SeccionID', 'TurnoID', 'ModalidadID'
    )
    
    # 3. Captura de filtros adicionales desde la URL
    selected_nivel = request.GET.get('nivel')
    selected_seccion = request.GET.get('seccion')
    selected_turno = request.GET.get('turno')
    selected_modalidad = request.GET.get('modalidad')

    # Nombres para el encabezado del reporte
    nombre_nivel = "TODOS"
    nombre_seccion = "TODAS"
    nombre_turno = "TODOS"
    nombre_modalidad = "TODAS"

    # Aplicación de filtros de búsqueda (si el usuario eligió algo en los select)
    if selected_nivel:
        matriculas = matriculas.filter(NivelID__NivelID=selected_nivel)
        nivel_obj = Nivel.objects.filter(NivelID=selected_nivel).first()
        if nivel_obj: nombre_nivel = nivel_obj.NombreNivel

    if selected_seccion:
        matriculas = matriculas.filter(SeccionID__SeccionID=selected_seccion)
        seccion_obj = Seccion.objects.filter(SeccionID=selected_seccion).first()
        if seccion_obj: nombre_seccion = seccion_obj.NombreSeccion

    if selected_turno:
        matriculas = matriculas.filter(TurnoID__TurnoID=selected_turno)
        turno_obj = Turno.objects.filter(TurnoID=selected_turno).first()
        if turno_obj: nombre_turno = turno_obj.NombreTurno

    if selected_modalidad:
        matriculas = matriculas.filter(ModalidadID__ModalidadID=selected_modalidad)
        mod_obj = Modalidades.objects.filter(ModalidadID=selected_modalidad).first()
        if mod_obj: nombre_modalidad = mod_obj.NombreModalidad
    
    # 4. Manejo de columnas dinámicas desde el Modal
    columnas_param = request.GET.get('columnas')
    if columnas_param:
        lista_columnas = columnas_param.split(',')
    else:
        # Configuración por defecto si no se usa el modal
        lista_columnas = ['cod_mined', 'genero', 'nivel', 'seccion', 'turno']

    # 5. Contexto para el template de impresión
    context = {
        'titulo': 'Reporte de Matrícula Activa',
        'matriculas': matriculas.order_by('EstudianteID__NombreCompleto'),
        'cols': lista_columnas,
        'anio_reporte': anio_actual,
        'nombre_nivel': nombre_nivel,
        'nombre_seccion': nombre_seccion,
        'nombre_turno': nombre_turno,
        'nombre_modalidad': nombre_modalidad,
    }
    
    return render(request, 'reportes/imprimir_reporte_completo.html', context)


from openpyxl.styles import Font, PatternFill

def exportar_estudiantes_excel_view(request):
    # 1. Obtener filtros de la URL
    selected_nivel_id = request.GET.get('nivel')
    selected_seccion_id = request.GET.get('seccion')
    
    # 2. Filtrar matrículas y estudiantes activos
    matriculas = Matricula.objects.filter(
        activo=True, 
        EstudianteID__activo=True 
    ).select_related(
        'EstudianteID', 'NivelID', 'SeccionID', 'TurnoID', 'ModalidadID'
    ).order_by('EstudianteID__NombreCompleto')

    # Aplicar filtros dinámicos
    if selected_nivel_id:
        matriculas = matriculas.filter(NivelID__NivelID=selected_nivel_id)
    if selected_seccion_id:
        matriculas = matriculas.filter(SeccionID__SeccionID=selected_seccion_id)
    # (Puedes repetir para turno y modalidad si lo deseas)

    # 3. Configuración de respuesta
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="reporte_estudiantes_activos.xlsx"'

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Estudiantes Activos'

    # --- NUEVO: INFORMACIÓN DEL REPORTE ---
    # Obtenemos los nombres de los filtros para mostrarlos en el Excel
    primer_registro = matriculas.first()
    txt_nivel = primer_registro.NivelID.NombreNivel if primer_registro and selected_nivel_id else "TODOS"
    txt_seccion = primer_registro.SeccionID.NombreSeccion if primer_registro and selected_seccion_id else "TODAS"

    sheet.append(['REPORTE DE ESTUDIANTES ACTIVOS'])
    sheet.append(['Nivel:', txt_nivel])
    sheet.append(['Sección:', txt_seccion])
    sheet.append([]) # Espacio en blanco

    # Estilos para encabezados
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

    # 4. Encabezados de la tabla (ahora en la fila 5)
    headers = [
        'N°', 'Codigo MINED', 'Nombre Completo', 'Genero', 'Fecha de Nacimiento',
        'Edad', 'Direccion', 'Nombre Madre', 'Cedula Madre', 'Telefono Madre',
        'Ocupacion Madre', 'Nombre Padre', 'Cedula Padre', 'Telefono Padre',
        'Ocupacion Padre', 'Nombre Tutor', 'Cedula Tutor', 'Telefono Tutor',
        'Ocupacion Tutor', 'Nivel', 'Seccion', 'Turno', 'Modalidad', 'Fecha Matricula',
        'Año Matricula'
    ]
    sheet.append(headers)

    # Aplicar estilo a los encabezados
    for cell in sheet[5]: # La fila 5 es donde quedaron los encabezados
        cell.font = header_font
        cell.fill = header_fill

    # 5. Llenado de datos
    for i, matricula in enumerate(matriculas, 1):
        est = matricula.EstudianteID
        row_data = [
            i, est.Codigo_MINED, est.NombreCompleto, est.Genero, est.FechaNacimiento,
            est.Edad, est.Direccion, est.NombreMadre, est.CedulaMadre, est.TelefonoMadre,
            est.OcupacionMadre, est.NombrePadre, est.CedulaPadre, est.TelefonoPadre,
            est.OcupacionPadre, est.NombreTutor, est.CedulaTutor, est.TelefonoTutor,
            est.OcupacionTutor, matricula.NivelID.NombreNivel, matricula.SeccionID.NombreSeccion,
            matricula.TurnoID.NombreTurno, matricula.ModalidadID.NombreModalidad,
            matricula.FechaMatricula, matricula.AñoMatricula,
        ]
        sheet.append(row_data)

    # 6. Ajuste de columnas
    for col in sheet.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except: pass
        sheet.column_dimensions[column].width = max_length + 2

    workbook.save(response)
    return response